"""Per-event guest attributes: definitions, answers and the summary."""
from datetime import date, datetime, timezone

import pytest

from rsvp_manager.extensions import db
from rsvp_manager.models import (
    Event, EventAttribute, EventAttributeOption, Guest, Invitation,
    InvitationAttributeValue,
)
from rsvp_manager.services import attribute_service


@pytest.fixture()
def hunt(test_app, user):
    e = Event(user_id=user, name="Boxing Day Hunt", event_type="Hunt",
              date=date(2026, 12, 26), date_created=date.today())
    db.session.add(e)
    db.session.commit()
    return e


def _guest(user, first, last="Test"):
    g = Guest(user_id=user, first_name=first, last_name=last, gender="Male",
              date_created=datetime.now(timezone.utc))
    db.session.add(g)
    db.session.commit()
    return g


def _invite(event, guest, status="Attending"):
    inv = Invitation(event_id=event.id, guest_id=guest.id, status=status)
    db.session.add(inv)
    db.session.commit()
    return inv


class TestDefinitions:
    def test_create_with_options(self, test_app, user, hunt):
        attr = attribute_service.create_attribute(hunt, "Hunting", ["Hunter", "Follower"])
        assert attr.name == "Hunting"
        assert [o.label for o in attr.options] == ["Hunter", "Follower"]

    def test_options_are_trimmed_deduped_and_ordered(self, test_app, user, hunt):
        attr = attribute_service.create_attribute(
            hunt, "Meal", ["  Lunch ", "Dinner", "lunch", "", "   ", "Lunch + Dinner"])
        assert [o.label for o in attr.options] == ["Lunch", "Dinner", "Lunch + Dinner"]

    def test_name_must_be_unique_per_event(self, test_app, user, hunt):
        attribute_service.create_attribute(hunt, "Hunting", ["Hunter"])
        with pytest.raises(Exception):
            attribute_service.create_attribute(hunt, "Hunting", ["Follower"])

    def test_same_name_allowed_on_a_different_event(self, test_app, user, hunt):
        other = Event(user_id=user, name="Second", event_type="Hunt",
                      date=date(2027, 1, 1), date_created=date.today())
        db.session.add(other)
        db.session.commit()
        attribute_service.create_attribute(hunt, "Hunting", ["Hunter"])
        attribute_service.create_attribute(other, "Hunting", ["Hunter"])
        assert len(attribute_service.get_attributes(other)) == 1

    def test_at_least_one_option_required(self, test_app, user, hunt):
        with pytest.raises(Exception):
            attribute_service.create_attribute(hunt, "Empty", ["  ", ""])

    def test_rename_option_keeps_guest_answers(self, test_app, user, hunt):
        attr = attribute_service.create_attribute(hunt, "Hunting", ["Hunter", "Follower"])
        inv = _invite(hunt, _guest(user, "Ann"))
        attribute_service.set_value(inv, attr, attr.options[1].id)

        attribute_service.rename_option(attr.options[1], "Non-hunter")

        assert inv.value_for(attr.id).label == "Non-hunter", (
            "renaming an option must not unset the guests who chose it"
        )

    def test_deleting_an_attribute_removes_its_answers(self, test_app, user, hunt):
        attr = attribute_service.create_attribute(hunt, "Hunting", ["Hunter"])
        inv = _invite(hunt, _guest(user, "Ann"))
        attribute_service.set_value(inv, attr, attr.options[0].id)
        attr_id = attr.id

        attribute_service.delete_attribute(attr)

        assert db.session.get(EventAttribute, attr_id) is None
        assert InvitationAttributeValue.query.filter_by(attribute_id=attr_id).count() == 0

    def test_deleting_the_event_removes_its_attributes(self, test_app, user, hunt):
        attr = attribute_service.create_attribute(hunt, "Hunting", ["Hunter"])
        attr_id, event_id = attr.id, hunt.id
        db.session.delete(hunt)
        db.session.commit()
        assert db.session.get(EventAttribute, attr_id) is None
        assert EventAttributeOption.query.filter_by(attribute_id=attr_id).count() == 0


class TestDefaultsPerEventType:
    def test_hunt_gets_the_hunting_attribute(self, test_app, user, hunt):
        created = attribute_service.apply_defaults_for_type(hunt)
        assert len(created) == 1
        assert created[0].name == "Hunting"
        assert [o.label for o in created[0].options] == ["Hunter", "Follower"]

    def test_other_types_get_nothing(self, test_app, user):
        party = Event(user_id=user, name="Party", event_type="Party",
                      date=date(2026, 6, 1), date_created=date.today())
        db.session.add(party)
        db.session.commit()
        assert attribute_service.apply_defaults_for_type(party) == []

    def test_applying_twice_does_not_duplicate(self, test_app, user, hunt):
        attribute_service.apply_defaults_for_type(hunt)
        attribute_service.apply_defaults_for_type(hunt)
        assert len(attribute_service.get_attributes(hunt)) == 1


class TestGuestAnswers:
    def test_set_change_and_clear(self, test_app, user, hunt):
        attr = attribute_service.create_attribute(hunt, "Hunting", ["Hunter", "Follower"])
        hunter, follower = attr.options
        inv = _invite(hunt, _guest(user, "Ann"))

        assert inv.value_for(attr.id) is None, "starts not set"

        attribute_service.set_value(inv, attr, hunter.id)
        assert inv.value_for(attr.id).label == "Hunter"

        attribute_service.set_value(inv, attr, follower.id)
        assert inv.value_for(attr.id).label == "Follower", "changing replaces, not duplicates"
        assert InvitationAttributeValue.query.filter_by(invitation_id=inv.id).count() == 1

        attribute_service.set_value(inv, attr, None)
        assert inv.value_for(attr.id) is None, "clearing returns to not set"

    def test_option_from_another_attribute_is_rejected(self, test_app, user, hunt):
        hunting = attribute_service.create_attribute(hunt, "Hunting", ["Hunter"])
        meal = attribute_service.create_attribute(hunt, "Meal", ["Lunch"])
        inv = _invite(hunt, _guest(user, "Ann"))
        with pytest.raises(Exception):
            attribute_service.set_value(inv, hunting, meal.options[0].id)

    def test_attribute_from_another_event_is_rejected(self, test_app, user, hunt):
        other = Event(user_id=user, name="Other", event_type="Hunt",
                      date=date(2027, 2, 2), date_created=date.today())
        db.session.add(other)
        db.session.commit()
        foreign = attribute_service.create_attribute(other, "Hunting", ["Hunter"])
        inv = _invite(hunt, _guest(user, "Ann"))
        with pytest.raises(Exception):
            attribute_service.set_value(inv, foreign, foreign.options[0].id)

    def test_answers_do_not_leak_between_events(self, test_app, user, hunt):
        """The same person can be a Hunter here and a Follower there."""
        other = Event(user_id=user, name="Other", event_type="Hunt",
                      date=date(2027, 3, 3), date_created=date.today())
        db.session.add(other)
        db.session.commit()
        guest = _guest(user, "Ann")
        a1 = attribute_service.create_attribute(hunt, "Hunting", ["Hunter", "Follower"])
        a2 = attribute_service.create_attribute(other, "Hunting", ["Hunter", "Follower"])
        i1, i2 = _invite(hunt, guest), _invite(other, guest)

        attribute_service.set_value(i1, a1, a1.options[0].id)
        attribute_service.set_value(i2, a2, a2.options[1].id)

        assert i1.value_for(a1.id).label == "Hunter"
        assert i2.value_for(a2.id).label == "Follower"

    def test_removing_a_guest_removes_their_answers(self, test_app, user, hunt):
        attr = attribute_service.create_attribute(hunt, "Hunting", ["Hunter"])
        inv = _invite(hunt, _guest(user, "Ann"))
        attribute_service.set_value(inv, attr, attr.options[0].id)
        inv_id = inv.id
        db.session.delete(inv)
        db.session.commit()
        assert InvitationAttributeValue.query.filter_by(invitation_id=inv_id).count() == 0


class TestBulkAssign:
    def test_sets_many_and_reports_only_what_changed(self, test_app, user, hunt):
        attr = attribute_service.create_attribute(hunt, "Hunting", ["Hunter", "Follower"])
        hunter = attr.options[0]
        invs = [_invite(hunt, _guest(user, f"G{i}")) for i in range(4)]

        changed = attribute_service.bulk_set_value(hunt, [i.id for i in invs], attr, hunter.id)
        assert len(changed) == 4

        again = attribute_service.bulk_set_value(hunt, [i.id for i in invs], attr, hunter.id)
        assert again == [], "setting the same value again changes nothing"

        cleared = attribute_service.bulk_set_value(hunt, [i.id for i in invs], attr, None)
        assert len(cleared) == 4

    def test_ignores_invitations_from_another_event(self, test_app, user, hunt):
        other = Event(user_id=user, name="Other", event_type="Hunt",
                      date=date(2027, 4, 4), date_created=date.today())
        db.session.add(other)
        db.session.commit()
        attr = attribute_service.create_attribute(hunt, "Hunting", ["Hunter"])
        mine = _invite(hunt, _guest(user, "Mine"))
        theirs = _invite(other, _guest(user, "Theirs"))

        changed = attribute_service.bulk_set_value(
            hunt, [mine.id, theirs.id], attr, attr.options[0].id)

        assert changed == [mine.id]
        assert theirs.value_for(attr.id) is None


class TestSummary:
    def test_counts_by_option_and_status(self, test_app, user, hunt):
        attr = attribute_service.create_attribute(hunt, "Hunting", ["Hunter", "Follower"])
        hunter, follower = attr.options

        for name, status, option in [
            ("A", "Attending", hunter), ("B", "Attending", hunter),
            ("C", "Pending", hunter), ("D", "Attending", follower),
            ("E", "Declined", follower), ("F", "Attending", None),
            ("G", "Not Sent", hunter),
        ]:
            inv = _invite(hunt, _guest(user, name), status)
            if option:
                attribute_service.set_value(inv, attr, option.id)

        summary = attribute_service.summarize(hunt)
        assert len(summary) == 1
        rows = {r["label"]: r for r in summary[0]["rows"]}

        assert rows["Hunter"]["attending"] == 2
        assert rows["Hunter"]["pending"] == 1
        assert rows["Hunter"]["invited"] == 3, "'Not Sent' is excluded from Invited"
        assert rows["Follower"]["attending"] == 1
        assert rows["Follower"]["declined"] == 1
        assert rows["Not set"]["attending"] == 1
        assert summary[0]["totals"]["attending"] == 4

    def test_rows_add_up_to_the_total(self, test_app, user, hunt):
        attr = attribute_service.create_attribute(hunt, "Hunting", ["Hunter", "Follower"])
        for i, opt in enumerate([attr.options[0], attr.options[1], None, attr.options[0]]):
            inv = _invite(hunt, _guest(user, f"G{i}"))
            if opt:
                attribute_service.set_value(inv, attr, opt.id)

        summary = attribute_service.summarize(hunt)[0]
        for key in ("attending", "pending", "declined", "invited"):
            assert sum(r[key] for r in summary["rows"]) == summary["totals"][key], (
                f"the {key} column must add up to the total row"
            )

    def test_not_set_row_is_hidden_when_everyone_is_set(self, test_app, user, hunt):
        attr = attribute_service.create_attribute(hunt, "Hunting", ["Hunter"])
        inv = _invite(hunt, _guest(user, "A"))
        attribute_service.set_value(inv, attr, attr.options[0].id)
        labels = [r["label"] for r in attribute_service.summarize(hunt)[0]["rows"]]
        assert "Not set" not in labels

    def test_trashed_guests_are_excluded(self, test_app, user, hunt):
        attr = attribute_service.create_attribute(hunt, "Hunting", ["Hunter"])
        guest = _guest(user, "Gone")
        inv = _invite(hunt, guest)
        attribute_service.set_value(inv, attr, attr.options[0].id)

        guest.deleted_at = datetime.now(timezone.utc)
        db.session.commit()

        rows = {r["label"]: r for r in attribute_service.summarize(hunt)[0]["rows"]}
        assert rows["Hunter"]["attending"] == 0

    def test_no_attributes_means_no_summary(self, test_app, user, hunt):
        assert attribute_service.summarize(hunt) == []


# ── API ─────────────────────────────────────────────────────────────────────

class TestAttributeApi:
    def _hunt_for(self, client, user):
        e = Event(user_id=user, name="API Hunt", event_type="Hunt",
                  date=date(2026, 12, 26), date_created=date.today())
        db.session.add(e)
        db.session.commit()
        return e

    def test_create_list_update_delete(self, logged_in_client, test_app, user):
        event = self._hunt_for(logged_in_client, user)
        csrf = _csrf(logged_in_client)

        r = logged_in_client.post(f"/api/v1/events/{event.id}/attributes",
                                  json={"name": "Hunting", "options": ["Hunter", "Follower"]},
                                  headers={"X-CSRFToken": csrf})
        assert r.status_code == 201, r.data
        attr = r.get_json()["data"]
        assert [o["label"] for o in attr["options"]] == ["Hunter", "Follower"]

        r = logged_in_client.get(f"/api/v1/events/{event.id}/attributes")
        assert len(r.get_json()["data"]) == 1

        r = logged_in_client.put(f"/api/v1/events/{event.id}/attributes/{attr['id']}",
                                 json={"name": "Role"}, headers={"X-CSRFToken": csrf})
        assert r.get_json()["data"]["name"] == "Role"

        r = logged_in_client.delete(f"/api/v1/events/{event.id}/attributes/{attr['id']}",
                                    headers={"X-CSRFToken": csrf})
        assert r.status_code == 204
        assert logged_in_client.get(f"/api/v1/events/{event.id}/attributes").get_json()["data"] == []

    def test_set_and_clear_a_guest_answer(self, logged_in_client, test_app, user):
        event = self._hunt_for(logged_in_client, user)
        csrf = _csrf(logged_in_client)
        attr = attribute_service.create_attribute(event, "Hunting", ["Hunter", "Follower"])
        inv = _invite(event, _guest(user, "Ann"))

        r = logged_in_client.put(
            f"/api/v1/invitations/{inv.id}/attributes/{attr.id}",
            json={"option_id": attr.options[0].id}, headers={"X-CSRFToken": csrf})
        assert r.status_code == 200
        assert r.get_json()["data"]["label"] == "Hunter"

        r = logged_in_client.put(
            f"/api/v1/invitations/{inv.id}/attributes/{attr.id}",
            json={"option_id": None}, headers={"X-CSRFToken": csrf})
        assert r.get_json()["data"]["option_id"] is None

    def test_bulk_set(self, logged_in_client, test_app, user):
        event = self._hunt_for(logged_in_client, user)
        csrf = _csrf(logged_in_client)
        attr = attribute_service.create_attribute(event, "Hunting", ["Hunter"])
        invs = [_invite(event, _guest(user, f"G{i}")) for i in range(3)]

        r = logged_in_client.post(
            f"/api/v1/events/{event.id}/attributes/{attr.id}/bulk",
            json={"invitation_ids": [i.id for i in invs], "option_id": attr.options[0].id},
            headers={"X-CSRFToken": csrf})
        assert r.get_json()["data"]["changed"] == 3

    def test_summary_endpoint(self, logged_in_client, test_app, user):
        event = self._hunt_for(logged_in_client, user)
        attr = attribute_service.create_attribute(event, "Hunting", ["Hunter"])
        inv = _invite(event, _guest(user, "Ann"))
        attribute_service.set_value(inv, attr, attr.options[0].id)

        data = logged_in_client.get(f"/api/v1/events/{event.id}/attributes/summary").get_json()["data"]
        assert data[0]["attribute"]["name"] == "Hunting"
        assert data[0]["rows"][0]["attending"] == 1

    def test_another_users_event_is_refused(self, logged_in_client, test_app, user, user2):
        """Same cross-tenant checks the audit applied to every other endpoint."""
        theirs = Event(user_id=user2, name="Theirs", event_type="Hunt",
                       date=date(2026, 12, 26), date_created=date.today())
        db.session.add(theirs)
        db.session.commit()
        attr = attribute_service.create_attribute(theirs, "Hunting", ["Hunter"])
        csrf = _csrf(logged_in_client)
        h = {"X-CSRFToken": csrf}

        assert logged_in_client.get(f"/api/v1/events/{theirs.id}/attributes").status_code in (403, 404)
        assert logged_in_client.post(f"/api/v1/events/{theirs.id}/attributes",
                                     json={"name": "X", "options": ["a"]}, headers=h
                                     ).status_code in (403, 404)
        assert logged_in_client.delete(f"/api/v1/events/{theirs.id}/attributes/{attr.id}",
                                       headers=h).status_code in (403, 404)
        assert logged_in_client.get(
            f"/api/v1/events/{theirs.id}/attributes/summary").status_code in (403, 404)

    def test_anonymous_is_refused(self, client, test_app, user):
        event = Event(user_id=user, name="E", event_type="Hunt",
                      date=date(2026, 12, 26), date_created=date.today())
        db.session.add(event)
        db.session.commit()
        assert client.get(f"/api/v1/events/{event.id}/attributes").status_code == 401


def _csrf(client):
    import re
    html = client.get("/settings").data.decode()
    return re.search(r'name="csrf_token"[^>]*value="([^"]+)"', html).group(1)


class TestHuntEventsStartWithTheHuntingAttribute:
    def test_created_hunt_is_prefilled(self, test_app, user):
        from rsvp_manager.services import event_service

        event = event_service.create_event(user, {
            "name": "Boxing Day", "event_type": "Hunt", "date": "2026-12-26",
        })
        attrs = attribute_service.get_attributes(event)
        assert [a.name for a in attrs] == ["Hunting"]
        assert [o.label for o in attrs[0].options] == ["Hunter", "Follower"]

    def test_created_party_is_not(self, test_app, user):
        from rsvp_manager.services import event_service

        event = event_service.create_event(user, {
            "name": "Summer Party", "event_type": "Party", "date": "2026-06-01",
        })
        assert attribute_service.get_attributes(event) == []


class TestExportIncludesAttributes:
    def test_one_column_per_attribute_with_the_guest_answer(self, logged_in_client, test_app, user):
        import io
        import openpyxl

        event = Event(user_id=user, name="Hunt Export", event_type="Hunt",
                      date=date(2026, 12, 26), date_created=date.today())
        db.session.add(event)
        db.session.commit()
        hunting = attribute_service.create_attribute(event, "Hunting", ["Hunter", "Follower"])
        meal = attribute_service.create_attribute(event, "Meal", ["Lunch", "Dinner"])

        set_guest = _invite(event, _guest(user, "Set"))
        _invite(event, _guest(user, "Unset"))
        attribute_service.set_value(set_guest, hunting, hunting.options[0].id)

        resp = logged_in_client.get(f"/export/event/{event.id}")
        assert resp.status_code == 200
        ws = openpyxl.load_workbook(io.BytesIO(resp.data)).active
        header = [c.value for c in ws[1]]
        assert header[-2:] == ["Hunting", "Meal"], header

        rows = {r[1]: r for r in ws.iter_rows(min_row=2, values_only=True)}
        assert rows["Set"][-2] == "Hunter"
        # openpyxl reads an empty cell back as None
        assert rows["Set"][-1] in (None, ""), "an unanswered attribute exports as blank"
        assert rows["Unset"][-2] in (None, "")

    def test_export_query_count_stays_flat(self, logged_in_client, test_app, user):
        from sqlalchemy import event as sa_event
        from sqlalchemy.engine import Engine

        event = Event(user_id=user, name="Counted", event_type="Hunt",
                      date=date(2026, 12, 26), date_created=date.today())
        db.session.add(event)
        db.session.commit()
        attr = attribute_service.create_attribute(event, "Hunting", ["Hunter"])

        def count(n_guests):
            for i in range(n_guests):
                inv = _invite(event, _guest(user, f"X{i}{n_guests}"))
                attribute_service.set_value(inv, attr, attr.options[0].id)
            seen = []

            def before(conn, cursor, statement, params, context, executemany):
                seen.append(statement)

            sa_event.listen(Engine, "before_cursor_execute", before)
            try:
                assert logged_in_client.get(f"/export/event/{event.id}").status_code == 200
            finally:
                sa_event.remove(Engine, "before_cursor_execute", before)
            return len(seen)

        few, many = count(2), count(20)
        assert many <= few + 2, (
            f"export queries grew with guest count ({few} -> {many}); the attribute "
            f"answers are not being eager-loaded"
        )


class TestAttributesFromTheCreateEventForm:
    """The new-event form sends parallel attribute_name / attribute_options fields."""

    def _form(self, pairs, event_type="Hunt"):
        from werkzeug.datastructures import MultiDict

        md = MultiDict([("name", "Formed"), ("event_type", event_type), ("date", "2026-12-26")])
        for name, options in pairs:
            md.add("attribute_name", name)
            md.add("attribute_options", options)
        return md

    def test_creates_what_the_form_describes(self, test_app, user):
        from rsvp_manager.services import event_service

        event = event_service.create_event(user, self._form([
            ("Hunting", "Hunter\nFollower"),
            ("Participating to", "Lunch\nDinner\nBoth"),
        ]))
        attrs = attribute_service.get_attributes(event)
        assert [a.name for a in attrs] == ["Hunting", "Participating to"]
        assert [o.label for o in attrs[1].options] == ["Lunch", "Dinner", "Both"]

    def test_user_can_remove_the_prefilled_default(self, test_app, user):
        """Submitting a Hunt with no attribute blocks means they wanted none."""
        from werkzeug.datastructures import MultiDict
        from rsvp_manager.services import event_service

        md = MultiDict([("name", "Bare"), ("event_type", "Hunt"), ("date", "2026-12-26"),
                        ("attribute_name", ""), ("attribute_options", "")])
        event = event_service.create_event(user, md)
        assert attribute_service.get_attributes(event) == [], (
            "clearing the blocks must not be overridden by the type default"
        )

    def test_blank_and_incomplete_blocks_are_skipped(self, test_app, user):
        from rsvp_manager.services import event_service

        event = event_service.create_event(user, self._form([
            ("Hunting", "Hunter\nFollower"),
            ("", "orphan answers"),
            ("Named but empty", "   \n  "),
        ]))
        assert [a.name for a in attribute_service.get_attributes(event)] == ["Hunting"]

    def test_a_form_without_attribute_fields_still_gets_the_defaults(self, test_app, user):
        """Older clients and API callers must not silently lose the Hunt default."""
        from rsvp_manager.services import event_service

        event = event_service.create_event(user, {
            "name": "Plain", "event_type": "Hunt", "date": "2026-12-26",
        })
        assert [a.name for a in attribute_service.get_attributes(event)] == ["Hunting"]


# ── Undo for "Remove from Event" ────────────────────────────────────────────

class TestUndoRemoveFromEvent:
    """Removing a guest is a hard delete the trash does not cover, so the
    snapshot the delete returns is the only way back."""

    def _setup(self, user, with_attribute=True, with_seat=False):
        from rsvp_manager.models import SeatAssignment, SeatingTable

        event = Event(user_id=user, name="Undo Hunt", event_type="Hunt",
                      date=date(2026, 12, 26), date_created=date.today())
        db.session.add(event)
        db.session.commit()
        attr = attribute_service.create_attribute(event, "Hunting", ["Hunter", "Follower"]) \
            if with_attribute else None
        guest = _guest(user, "Ann")
        inv = Invitation(event_id=event.id, guest_id=guest.id, status="Attending",
                         notes="bring boots", date_invited=date(2026, 11, 1))
        db.session.add(inv)
        db.session.commit()
        if attr:
            attribute_service.set_value(inv, attr, attr.options[0].id)
        if with_seat:
            table = SeatingTable(event_id=event.id, table_number=1, capacity=8, shape="round")
            db.session.add(table)
            db.session.commit()
            db.session.add(SeatAssignment(table_id=table.id, invitation_id=inv.id, seat_position=2))
            db.session.commit()
        return event, guest, inv, attr

    def test_snapshot_then_restore_brings_everything_back(self, test_app, user):
        from rsvp_manager.services import invitation_service

        event, guest, inv, attr = self._setup(user)
        snap = invitation_service.snapshot_invitation(inv)
        invitation_service.remove_invitation(inv)
        assert Invitation.query.filter_by(event_id=event.id).count() == 0

        restored = invitation_service.restore_invitations(event, [snap], user)

        assert len(restored) == 1
        back = restored[0]
        assert back.guest_id == guest.id
        assert back.status == "Attending"
        assert back.notes == "bring boots"
        assert back.date_invited == date(2026, 11, 1)
        assert back.value_for(attr.id).label == "Hunter", "their attribute answer comes back too"

    def test_seat_is_restored_when_the_chair_is_still_free(self, test_app, user):
        from rsvp_manager.models import SeatAssignment
        from rsvp_manager.services import invitation_service

        event, guest, inv, _ = self._setup(user, with_seat=True)
        snap = invitation_service.snapshot_invitation(inv)
        invitation_service.remove_invitation(inv)
        assert SeatAssignment.query.count() == 0

        restored = invitation_service.restore_invitations(event, [snap], user)
        seat = SeatAssignment.query.filter_by(invitation_id=restored[0].id).first()
        assert seat is not None and seat.seat_position == 2

    def test_undoing_twice_does_not_duplicate_the_guest(self, test_app, user):
        from rsvp_manager.services import invitation_service

        event, guest, inv, _ = self._setup(user)
        snap = invitation_service.snapshot_invitation(inv)
        invitation_service.remove_invitation(inv)

        invitation_service.restore_invitations(event, [snap], user)
        second = invitation_service.restore_invitations(event, [snap], user)

        assert second == [], "a guest already back on the event is skipped"
        assert Invitation.query.filter_by(event_id=event.id, guest_id=guest.id).count() == 1

    def test_answers_for_a_deleted_attribute_are_dropped_not_fatal(self, test_app, user):
        from rsvp_manager.services import invitation_service

        event, guest, inv, attr = self._setup(user)
        snap = invitation_service.snapshot_invitation(inv)
        invitation_service.remove_invitation(inv)
        attribute_service.delete_attribute(attr)

        restored = invitation_service.restore_invitations(event, [snap], user)
        assert len(restored) == 1, "undo still works after the attribute was removed"

    def test_another_users_guest_is_not_restored(self, test_app, user, user2):
        from rsvp_manager.services import invitation_service

        event, guest, inv, _ = self._setup(user)
        snap = invitation_service.snapshot_invitation(inv)
        invitation_service.remove_invitation(inv)
        snap["guest_id"] = _guest(user2, "Theirs").id

        assert invitation_service.restore_invitations(event, [snap], user) == []

    def test_delete_endpoint_returns_a_snapshot(self, logged_in_client, test_app, user):
        event, guest, inv, attr = self._setup(user)
        csrf = _csrf(logged_in_client)

        resp = logged_in_client.delete(f"/api/v1/invitations/{inv.id}",
                                       headers={"X-CSRFToken": csrf})
        assert resp.status_code == 200
        snap = resp.get_json()["data"]["snapshot"]
        assert snap["guest_id"] == guest.id
        assert snap["status"] == "Attending"
        assert snap["attributes"][0]["option_id"] == attr.options[0].id

        resp = logged_in_client.post(f"/api/v1/events/{event.id}/invitations/restore",
                                     json={"snapshots": [snap]},
                                     headers={"X-CSRFToken": csrf})
        assert resp.status_code == 200
        assert resp.get_json()["data"]["restored"] == 1

    def test_restore_on_another_users_event_is_refused(self, logged_in_client, test_app, user, user2):
        theirs = Event(user_id=user2, name="Theirs", event_type="Hunt",
                       date=date(2026, 12, 26), date_created=date.today())
        db.session.add(theirs)
        db.session.commit()
        resp = logged_in_client.post(f"/api/v1/events/{theirs.id}/invitations/restore",
                                     json={"snapshots": []},
                                     headers={"X-CSRFToken": _csrf(logged_in_client)})
        assert resp.status_code in (403, 404)
